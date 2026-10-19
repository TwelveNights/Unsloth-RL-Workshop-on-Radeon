#!/usr/bin/env python3
"""Adds the SFT mechanics and early-finisher challenge slides to the Studio deck.

Run from the repo root: python3 scripts/update_deck_sft_challenges.py
New slides copy the layout of existing ones (background, cards, colours) and only
change their text, so they match the rest of the deck. Re-export the PDF after.
"""
import copy
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn

DECK = "docs/AMD-Unsloth-Studio-Fine-tuning.pptx"
p = Presentation(DECK)
ids = p.slides._sldIdLst


def shape(slide, sid):
    return next(s for s in slide.shapes if s.shape_id == sid)


def runs(slide, sid, *texts, para=0):
    """Sets a paragraph's runs in order; leftover runs are emptied, so each keeps its formatting."""
    r = shape(slide, sid).text_frame.paragraphs[para].runs
    assert len(texts) <= len(r), (sid, para, [x.text for x in r])
    for i, run in enumerate(r):
        run.text = texts[i] if i < len(texts) else ""


def lines(slide, sid, texts):
    """Sets a box to one paragraph per text, copying the first paragraph's formatting."""
    tf = shape(slide, sid).text_frame
    paras = tf.paragraphs
    template = copy.deepcopy(paras[0]._p)
    for extra in paras[1:]:
        extra._p.getparent().remove(extra._p)
    for i, text in enumerate(texts):
        if i:
            tf._txBody.append(copy.deepcopy(template))
        runs(slide, sid, text, para=i)


def bullets(slide, sid, texts):
    """Bulleted box whose paragraphs are runs ['•  ', text, ...]."""
    tf = shape(slide, sid).text_frame
    template = copy.deepcopy(tf.paragraphs[0]._p)
    for extra in tf.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    for i, text in enumerate(texts):
        if i:
            tf._txBody.append(copy.deepcopy(template))
        runs(slide, sid, "•  ", text, para=i)


def duplicate(src_number, position):
    """Copies slide src_number (1-based, current order) to 1-based position."""
    src = p.slides[src_number - 1]
    new = p.slides.add_slide(src.slide_layout)
    # Copy into the new slide's own elements: python-pptx caches its shape tree.
    cSld, tree = new._element.cSld, new._element.cSld.spTree
    for child in list(tree):
        tree.remove(child)
    for child in src._element.cSld.spTree:
        tree.append(copy.deepcopy(child))
    if src._element.cSld.bg is not None:
        cSld.insert(0, copy.deepcopy(src._element.cSld.bg))
    # Re-link the background image (and any pictures) to the new slide.
    for el in new._element.iter():
        rid = el.get(qn("r:embed"))
        if rid:
            el.set(qn("r:embed"), new.part.relate_to(src.part.related_part(rid), RT.IMAGE))
    el = ids[-1]
    ids.remove(el)
    ids.insert(position - 1, el)
    return new


def remove_shapes(slide, top_from):
    """Removes every shape whose top edge is at or below top_from (EMU)."""
    for s in list(slide.shapes):
        if s.top >= top_from and not s.is_placeholder:
            s._element.getparent().remove(s._element)


# --- Slide 7: the adapter in this workshop is ~1% of the model, not 0.1%.
s7 = p.slides[6]
runs(s7, 16, "~1%")
runs(s7, 17, "of weights trained today (23M)")

# --- Slide 8: show the chat format attendees actually upload, not Alpaca.
s8 = p.slides[7]
runs(s8, 10, '{"messages": [', para=0)
runs(s8, 10, '  {"role": "user", "content": "<support message>"},', para=1)
runs(s8, 10, '  {"role": "assistant", "content": "<ticket JSON>"}]}', para=2)

# --- New 9: one SFT training step, on slide 6's four-card layout.
a = duplicate(6, 9)
runs(a, 2, "How SFT learns: one training step")
runs(a, 3, "Each step shows the model a few examples and nudges it toward ", "your", " answers, ", "token by token", ".")
for card, (title, body) in zip((7, 12, 17, 22), [
    ("Tokenize", "The chat template turns the message and its ticket into one sequence of tokens."),
    ("Predict", "At every position the model guesses the next token, just as when it chats."),
    ("Measure", "Cross-entropy loss: how unlikely it found the real next token. Lower is better."),
    ("Update", "Backpropagation adjusts only the LoRA adapter, making the right tokens likelier."),
]):
    runs(a, card, title)
    runs(a, card + 1, body)
runs(a, 26, "One step is 8 examples. ",
     "30 steps show the model 240 tickets, about a quarter of the data. ",
     "The test tickets are never trained on, which is why the notebook scores on them.")

# --- New 10: reading a training run, on slide 7's two-column layout.
b = duplicate(7, 10)
runs(b, 2, "Reading a training run")
runs(b, 3, "The settings Studio used today, and what the loss chart tells you.")
runs(b, 6, "Knobs")
runs(b, 7, "Studio's defaults for this run")
bullets(b, 8, ["Learning rate 2e-4: the size of each update. It warms up for 5 steps, then decays",
               "Batch of 8: 2 examples at a time, 4 accumulated per update",
               "LoRA rank 16: how much the adapter can learn"])
runs(b, 11, "Loss")
runs(b, 12, "How wrong the next-token guesses are")
bullets(b, 13, ["Falls fast: the model is picking up the format",
                "Flat near zero: little left to learn from this data",
                "Only held-out tests show whether it generalizes or memorized"])
runs(b, 16, "240")
runs(b, 17, "examples seen in 30 steps")
runs(b, 20, "~1%")
runs(b, 21, "of the weights trained (23M)")
runs(b, 24, "~4 min")
runs(b, 25, "for 30 steps on a Ryzen AI Max+")

# Learn more is now slide 20 (18 + the two above). The challenges go before it.
# --- New 20: stretch challenges, on the troubleshooting layout (first three rows).
c = duplicate(19, 20)
runs(c, 2, "Finished early? Stretch challenges")
runs(c, 3, "CHALLENGE")
runs(c, 4, "HOW")
for title_id, body_id, title, body in [
    (7, 8, "Add a field", "Add priority (urgent when the message says ASAP) to the data generator and the scorer's FIELDS, then retrain."),
    (11, 12, "Teach it your own task", "Write 5 examples by hand, let a bigger model write 50 more, train 30 steps. Ideas on the next slide."),
    (15, 16, "Take it with you", "Export as GGUF (Q4_K_M) and run it in llama.cpp, LM Studio, Ollama or Lemonade. Compare answers and speed."),
]:
    runs(c, title_id, title)
    runs(c, body_id, body)
remove_shapes(c, shape(c, 19).top - 320000)   # rows 4-6 (19 is row 4's title)

# --- New 21: your own task, on the SFT/GRPO two-column layout.
d = duplicate(8, 21)
runs(d, 2, "Teach it your own task")
runs(d, 3, "Any task with one right answer per input works the same way as today's tickets.")
tags = [s for s in d.shapes if s.has_text_frame and s.text_frame.text in ("SFT", "GRPO")]
for s, text in zip(sorted(tags, key=lambda s: s.left), ("IDEAS", "STEPS")):
    s.text_frame.paragraphs[0].runs[0].text = text
runs(d, 7, "Pick something small")
bullets(d, 8, ["One right answer per input, so you can check it",
               "A short, fixed output, like today's seven fields"])
lines(d, 10, ["meeting notes  -> {owner, task, due}",
              "receipt text   -> {merchant, date, total}",
              "bug report     -> {component, severity}",
              "job ad         -> {title, remote, salary}",
              "casual email   -> a polite, formal reply"])
runs(d, 14, "Write a few, generate the rest")
bullets(d, 15, ["Your hand-written examples set the format and the bar",
                "A bigger model adds variety: Qwen3.6-35B in Lemonade"])
runs(d, 17, "Write  ", "5 examples by hand, covering the cases you care about")
runs(d, 19, "Generate  ", "50 more: paste yours and ask for more like them")
runs(d, 21, "Check  ", "each one, keep 10 back for testing, train 30 steps")

p.save(DECK)
print("saved", DECK, len(p.slides), "slides")
